import hashlib
import io
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas
from docx import Document

from app import app

VERSION = "c0ef5a54-0fd8-41aa-b89e-d3cae8138523"


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("DOCLING_INTERNAL_TOKEN", "test-internal-service-secret")
    with TestClient(app) as client:
        yield client


def upload(client, content=b"Hello Ticknova", name="guide.txt", mime="text/plain", **overrides):
    data = {"documentVersionId": VERSION, "sha256": hashlib.sha256(content).hexdigest(), "originalFileName": name}
    data.update(overrides.pop("data", {}))
    headers = overrides.pop("headers", {"Authorization": "Bearer test-internal-service-secret"})
    return client.post("/internal/v1/convert", data=data, files={"sourceFile": (name, content, mime)}, headers=headers, **overrides)


@pytest.mark.parametrize("headers", [{}, {"Authorization": "Bearer wrong"}])
def test_auth_rejects_before_body_consumption(client, headers):
    response = client.post("/internal/v1/convert", content=b"malformed multipart", headers=headers)
    assert response.status_code == 401
    assert response.json() == {"errorCode": "UNAUTHORIZED"}


def test_missing_config_fails_closed(client, monkeypatch):
    monkeypatch.delenv("DOCLING_INTERNAL_TOKEN")
    assert upload(client).status_code == 503


@pytest.mark.parametrize("name,mime,body,code", [
    ("guide.exe", "application/octet-stream", b"MZ", "UNSUPPORTED_FORMAT"),
    ("../guide.txt", "text/plain", b"text", "INVALID_FILENAME"),
    ("C:\\guide.txt", "text/plain", b"text", "INVALID_FILENAME"),
    ("guide.pdf", "application/pdf", b"forged", "INVALID_DOCUMENT"),
    ("guide.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PKforged", "INVALID_DOCUMENT"),
    ("guide.txt", "application/pdf", b"text", "MIME_MISMATCH"),
    ("guide.txt", "text/plain", b"\xff\x00", "INVALID_UTF8"),
    ("guide.txt", "text/plain", b"%PDF-1.7 fake", "INVALID_DOCUMENT"),
])
def test_invalid_input_is_rejected(client, name, mime, body, code):
    response = upload(client, body, name, mime)
    assert response.status_code == 400
    assert response.json() == {"errorCode": code}


def test_hash_mismatch(client):
    response = upload(client, data={"sha256": "0" * 64})
    assert response.status_code == 400
    assert response.json() == {"errorCode": "HASH_MISMATCH"}


@pytest.mark.parametrize("data", [{"url": "http://example.com"}, {"path": "/etc/passwd"}, {"documentVersionId": "invalid"}])
def test_unknown_and_invalid_fields_rejected(client, data):
    response = upload(client, data=data)
    assert response.status_code == 400
    assert response.json() == {"errorCode": "INVALID_REQUEST"}


def test_file_size_limit(client):
    response = upload(client, b"x" * (20 * 1024 * 1024 + 1))
    assert response.status_code == 413
    assert response.json() == {"errorCode": "FILE_TOO_LARGE"}


def test_pdf_page_limit(client):
    stream = io.BytesIO()
    writer = PdfWriter()
    for _ in range(101):
        writer.add_blank_page(width=100, height=100)
    writer.write(stream)
    response = upload(client, stream.getvalue(), "guide.pdf", "application/pdf")
    assert response.status_code == 400
    assert response.json() == {"errorCode": "PDF_PAGE_LIMIT"}


def valid_document(kind):
    if kind == "pdf":
        stream = io.BytesIO()
        canvas = Canvas(stream)
        canvas.drawString(72, 720, "Hello Ticknova")
        canvas.save()
        return stream.getvalue(), "application/pdf"
    if kind == "docx":
        stream = io.BytesIO()
        document = Document()
        document.add_heading("Manual", 1)
        document.add_paragraph("Hello Ticknova")
        document.save(stream)
        return stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return (b"# Manual\n\nHello Ticknova" if kind == "md" else b"Hello Ticknova"), ("text/markdown" if kind == "md" else "text/plain")


@pytest.mark.parametrize("kind", ["pdf", "docx", "md", "txt"])
def test_real_docling_conversion(client, kind):
    body, mime = valid_document(kind)
    response = upload(client, body, "guide." + kind, mime)
    assert response.status_code == 200, response.text
    result = response.json()
    assert UUID(result["conversionId"])
    assert result["documentVersionId"] == VERSION
    assert result["sourceSha256"] == hashlib.sha256(body).hexdigest()
    assert result["status"] == "SUCCEEDED"
    assert "Hello Ticknova" in result["markdown"]
    assert any("Hello Ticknova" in page["text"] for page in result["pages"])
    assert all(page["pageNo"] >= 1 and isinstance(page["headingPath"], list) for page in result["pages"])
    assert "docling/2.142.1" in result["parserVersion"]
    assert isinstance(result["warnings"], list)
    if kind in {"docx", "md"}:
        assert any(page["headingPath"] == ["Manual"] for page in result["pages"])


def test_timeout_kills_worker_and_cleans_temp(client, monkeypatch, tmp_path):
    import app as service
    monkeypatch.setattr(service, "CONVERSION_TIMEOUT_SECONDS", 0.2)
    monkeypatch.setattr(service, "TEMP_ROOT", str(tmp_path))
    response = upload(client)
    assert response.status_code == 504
    assert response.json() == {"errorCode": "CONVERSION_TIMEOUT"}
    assert list(tmp_path.iterdir()) == []


def test_timeout_after_headers_never_sends_a_second_response_start(monkeypatch):
    import asyncio
    import app as service

    monkeypatch.setenv("DOCLING_INTERNAL_TOKEN", "test-internal-service-secret")
    monkeypatch.setattr(service, "CONVERSION_TIMEOUT_SECONDS", 0.01)
    sent = []

    async def stalled_response(scope, receive, send):
        await send({"type": "http.response.start", "status": 200, "headers": []})
        await asyncio.sleep(1)

    async def receive():
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(message):
        sent.append(message)

    boundary = service.InternalBoundary(stalled_response)
    scope = {"type": "http", "path": "/internal/v1/convert", "headers": [(b"authorization", b"Bearer test-internal-service-secret")]}
    with pytest.raises(TimeoutError):
        asyncio.run(boundary(scope, receive, send))
    assert [message["status"] for message in sent if message["type"] == "http.response.start"] == [200]
    assert boundary.active == 0


def test_worker_exit_racing_timeout_kill_preserves_504_and_cleanup(client, monkeypatch, tmp_path):
    import asyncio
    import app as service

    class ExitingProcess:
        returncode = None
        waits = 0
        killed = False

        async def wait(self):
            self.waits += 1
            if self.waits == 1:
                await asyncio.sleep(10)
            self.returncode = 0
            return 0

        def kill(self):
            self.killed = True
            raise ProcessLookupError("Worker exited between returncode check and kill")

    process = ExitingProcess()

    async def create_process(*args, **kwargs):
        return process

    monkeypatch.setattr(service.asyncio, "create_subprocess_exec", create_process)
    monkeypatch.setattr(service, "CONVERSION_TIMEOUT_SECONDS", 0.2)
    monkeypatch.setattr(service, "TEMP_ROOT", str(tmp_path))
    response = upload(client)
    assert response.status_code == 504
    assert response.json() == {"errorCode": "CONVERSION_TIMEOUT"}
    assert process.killed and process.waits == 2
    assert list(tmp_path.iterdir()) == []


def test_empty_pdf_does_not_succeed(client):
    stream = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(stream)
    response = upload(client, stream.getvalue(), "scan.pdf", "application/pdf")
    assert response.status_code == 422
    assert response.json() == {"errorCode": "NO_EXTRACTABLE_TEXT"}


def test_pdf_retains_physical_page_numbers(client):
    stream = io.BytesIO()
    canvas = Canvas(stream)
    canvas.drawString(72, 720, "Page one")
    canvas.showPage()
    canvas.drawString(72, 720, "Page two")
    canvas.save()
    response = upload(client, stream.getvalue(), "manual.pdf", "application/pdf")
    assert response.status_code == 200, response.text
    assert [(page["pageNo"], page["text"].strip()) for page in response.json()["pages"]] == [(1, "Page one"), (2, "Page two")]


def test_markdown_remote_and_local_images_are_never_fetched(client, tmp_path):
    secret = tmp_path / "private.txt"
    secret.write_text("PRIVATE_CONTENT_MUST_STAY_PRIVATE")
    content = f"# Manual\n\nHello Ticknova\n\n![remote](http://127.0.0.1:9/a.png)\n\n![local](file:///{secret.as_posix()})".encode()
    response = upload(client, content, "manual.md", "text/markdown")
    assert response.status_code == 200, response.text
    assert "PRIVATE_CONTENT_MUST_STAY_PRIVATE" not in response.text


def test_bounded_output_rejects_large_text(client):
    response = upload(client, b"word " * 500_000)
    assert response.status_code == 422
    assert response.json() == {"errorCode": "OUTPUT_TOO_LARGE"}


def test_malformed_multipart_is_sanitized(client):
    response = client.post("/internal/v1/convert", content=b"broken", headers={
        "Authorization": "Bearer test-internal-service-secret", "Content-Type": "multipart/form-data; boundary=x"})
    assert response.status_code == 400
    assert response.json() == {"errorCode": "INVALID_REQUEST"}


def test_total_chunked_body_limit_has_stable_code(client):
    response = upload(client, b"x" * (21 * 1024 * 1024))
    assert response.status_code == 413
    assert response.json() == {"errorCode": "FILE_TOO_LARGE"}


def test_stream_overflow_closes_multipart_spools(client, monkeypatch):
    import asyncio
    import httpx
    import starlette.formparsers as forms
    from app import app
    opened = []
    original = forms.SpooledTemporaryFile

    def track(*args, **kwargs):
        file = original(*args, **kwargs)
        opened.append(file)
        return file

    monkeypatch.setattr(forms, "SpooledTemporaryFile", track)
    request = httpx.Request("POST", "http://test/internal/v1/convert", files={"sourceFile": ("manual.txt", b"x" * (21 * 1024 * 1024), "text/plain")})
    body = request.read()
    chunks = iter(body[i:i+65536] for i in range(0, len(body), 65536))
    sent = []

    async def receive():
        chunk = next(chunks, b"")
        return {"type": "http.request", "body": chunk, "more_body": bool(chunk)}

    async def send(message):
        sent.append(message)

    scope = {"type": "http", "method": "POST", "path": "/internal/v1/convert", "query_string": b"", "headers": [
        (b"content-type", request.headers["content-type"].encode()),
        (b"authorization", b"Bearer test-internal-service-secret")], "server": ("test", 80), "client": ("test", 1), "scheme": "http", "http_version": "1.1"}
    asyncio.run(app(scope, receive, send))
    assert sent[0]["status"] == 413
    assert opened and all(file.closed for file in opened)


def test_upload_size_can_be_lowered(client, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_BYTES", "16")
    response = upload(client, b"x" * 17)
    assert response.status_code == 413
    assert response.json() == {"errorCode": "FILE_TOO_LARGE"}


def test_pdf_page_limit_can_be_lowered(client, monkeypatch):
    monkeypatch.setenv("MAX_PDF_PAGES", "1")
    stream = io.BytesIO()
    writer = PdfWriter()
    for _ in range(2):
        writer.add_blank_page(width=100, height=100)
    writer.write(stream)
    response = upload(client, stream.getvalue(), "manual.pdf", "application/pdf")
    assert response.status_code == 400
    assert response.json() == {"errorCode": "PDF_PAGE_LIMIT"}


def test_invalid_limits_fail_closed(client, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_BYTES", "not-an-integer")
    response = upload(client)
    assert response.status_code == 503
    assert response.json() == {"errorCode": "SERVICE_UNAVAILABLE"}


def test_missing_mime_is_rejected_without_details(client):
    response = upload(client, mime="")
    assert response.status_code == 400
    assert response.json() == {"errorCode": "MIME_MISMATCH"}


def test_txt_http_response_preserves_technical_text(client):
    text = "Use <host>, **stars**, `code`.\r\n# shell comment\r\nprintf '%s\\n' \"$X\"\r\n<config>value</config>\r\n"
    response = upload(client, text.encode("utf-8"))
    assert response.status_code == 200, response.text
    assert response.json()["markdown"] == text
    assert response.json()["pages"] == [{"pageNo": 1, "headingPath": [], "text": text}]


def test_heading_amplification_returns_small_stable_error(client):
    content = "## " + "P" * 16000 + "\n\n" + "".join(f"### Section {i}\n\nBody\n\n" for i in range(600))
    response = upload(client, content.encode("utf-8"), "headings.md", "text/markdown")
    assert response.status_code == 422
    assert response.json() == {"errorCode": "OUTPUT_TOO_LARGE"}
    assert len(response.content) < 100
