"""Content fidelity and bounded provenance regressions from independent review."""
import json
import socket

import pytest

from worker import InvalidDocument, convert


@pytest.fixture
def worker_convert(monkeypatch):
    # The disposable production worker disables sockets. Restore the original
    # transports after exercising that same worker code in this test process.
    monkeypatch.setattr(socket.socket, "connect", socket.socket.connect)
    monkeypatch.setattr(socket.socket, "connect_ex", socket.socket.connect_ex)
    monkeypatch.setattr(socket, "getaddrinfo", socket.getaddrinfo)
    return convert


@pytest.mark.parametrize("text", [
    "Use the <host> argument.\r\nSet x to **literal stars**.\r\n",
    '<?xml version="1.0"?>\n<config>\n  <host>localhost</host>\n</config>\n',
    "#!/bin/sh\n# important comment\nexport X=one\nprintf '%s\\n' \"$X\"\n",
    "Keep `backticks`, _underscores_, **stars**, & ampersands.\n\nNext paragraph.\n",
])
def test_txt_preserves_literal_text_and_lines(worker_convert, tmp_path, text):
    source = tmp_path / "literal.txt"
    source.write_bytes(text.encode("utf-8"))
    result = worker_convert(source)
    assert result["markdown"] == text
    assert result["pages"] == [{"pageNo": 1, "headingPath": [], "text": text}]


def test_markdown_title_remains_parent_of_nested_sections(worker_convert, tmp_path):
    source = tmp_path / "hierarchy.md"
    source.write_text("# Parent\n\nParent text\n\n## Child\n\nChild text\n\n### Grandchild\n\nGrandchild text\n", encoding="utf-8")
    result = worker_convert(source)
    assert [page["headingPath"] for page in result["pages"]] == [
        ["Parent"], ["Parent", "Child"], ["Parent", "Child", "Grandchild"]]
    assert all(page["pageNo"] == 1 for page in result["pages"])


def test_docx_heading_levels_are_not_offset(worker_convert, tmp_path):
    from docx import Document
    source = tmp_path / "hierarchy.docx"
    document = Document()
    for level, text in enumerate(["Parent", "Child", "Grandchild"], 1):
        document.add_heading(text, level)
        document.add_paragraph(text + " text")
    document.save(source)
    assert [page["headingPath"] for page in worker_convert(source)["pages"]] == [
        ["Parent"], ["Parent", "Child"], ["Parent", "Child", "Grandchild"]]


@pytest.mark.parametrize("heading,repetitions", [("P" * 16000, 600), ("中" * 16000, 200), ('"' * 8000, 600)],
                         ids=["ascii", "utf8", "json-escapes"])
def test_repeated_heading_metadata_is_bounded_before_return(worker_convert, tmp_path, heading, repetitions):
    source = tmp_path / "amplification.md"
    source.write_text("## " + heading + "\n\n" + "".join(
        f"### Section {index}\n\nBody {index}\n\n" for index in range(repetitions)), encoding="utf-8")
    with pytest.raises(InvalidDocument, match="^OUTPUT_TOO_LARGE$"):
        worker_convert(source)


def test_ordinary_many_sections_remain_valid(worker_convert, tmp_path):
    source = tmp_path / "sections.md"
    source.write_text("# Parent\n\n" + "".join(
        f"## Section {index}\n\nBody {index}\n\n" for index in range(600)), encoding="utf-8")
    result = worker_convert(source)
    assert len(result["pages"]) == 601
    assert result["pages"][-1]["headingPath"] == ["Parent", "Section 599"]
    assert len(json.dumps(result, ensure_ascii=False).encode("utf-8")) < 8 * 1024 * 1024
